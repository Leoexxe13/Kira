import unittest
import json
import os
import tempfile
from unittest.mock import MagicMock
from actions.document_maker import handle_action as dm_handle

class TestDeliverableSemantics(unittest.TestCase):
    def test_semantic_xlsx_creation(self):
        """
        Prueba que el Deliverable Planner reciba metadata semántica
        equivalente a una intención tabular numérica y use openpyxl deterministicamente.
        """
        # Simulamos que Gemini dedujo el formato correcto y construyó los format_options.
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        
        args = {
            "action": "create_xlsx",
            "file_path": path,
            "title": "Reporte de Ventas",
            "content": [
                ["Producto", "Cantidad", "Precio Unitario", "Total"],
                ["Producto A", 10, 15.5, ""],
                ["Producto B", 5, 20.0, ""],
                ["Producto C", 8, 12.0, ""]
            ],
            "format_options": {
                "calculated_columns": {"D": "=B{row}*C{row}"},
                "totals_row": True,
                "currency_cols": ["C", "D"]
            }
        }
        
        res = dm_handle(args)
        self.assertIn("DELIVERABLE_CREATED", res)
        
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        
        # Verify headers
        self.assertEqual(ws["A1"].value, "Producto")
        self.assertEqual(ws["D1"].value, "Total")
        
        # Verify values and deterministic formulas
        self.assertEqual(ws["B2"].value, 10)
        self.assertEqual(ws["C2"].value, 15.5)
        self.assertEqual(ws["D2"].value, "=B2*C2") # Formula injected!
        self.assertEqual(ws["D3"].value, "=B3*C3")
        self.assertEqual(ws["D4"].value, "=B4*C4")
        
        # Totals row was automatically added at row 5
        self.assertEqual(ws["A5"].value, "TOTAL")
        self.assertEqual(ws["C5"].value, "=SUM(C2:C4)")
        self.assertEqual(ws["D5"].value, "=SUM(D2:D4)")
        
        # Format check (currency)
        self.assertEqual(ws["C2"].number_format, "$#,##0.00")
        self.assertEqual(ws["D2"].number_format, "$#,##0.00")
        self.assertEqual(ws["C5"].number_format, "$#,##0.00")
        self.assertEqual(ws["D5"].number_format, "$#,##0.00")
        
        wb.close()
        os.remove(path)

if __name__ == '__main__':
    unittest.main()

    def test_semantic_format_choices(self):
        """
        Pruebas simuladas de elección semántica donde verificamos que document_maker
        puede manejar los distintos tipos de formato. La inteligencia recae en Gemini,
        aquí verificamos que los drivers de KIRA los instancian bien.
        """
        # DOCX para narrativas
        args_docx = {
            "action": "create_docx",
            "file_path": tempfile.mktemp(suffix=".docx"),
            "title": "Documento Formal",
            "content": [{"type": "paragraph", "text": "Narrativa formal de prueba."}]
        }
        res_docx = dm_handle(args_docx)
        self.assertIn("DELIVERABLE_CREATED", res_docx)
        os.remove(args_docx["file_path"])
        
        # PPTX para exposición
        args_pptx = {
            "action": "create_pptx",
            "file_path": tempfile.mktemp(suffix=".pptx"),
            "title": "Presentación",
            "content": [{"layout": "content", "title": "Slide 1", "content": ["Punto 1"]}]
        }
        res_pptx = dm_handle(args_pptx)
        self.assertIn("DELIVERABLE_CREATED", res_pptx)
        os.remove(args_pptx["file_path"])
